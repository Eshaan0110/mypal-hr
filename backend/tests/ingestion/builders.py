"""Build small PDF and DOCX files in memory, so each parser test controls its exact input."""

import io

from docx import Document
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    ListFlowable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Table,
)

# (kind, content) with kind in: title, h1, h2, p, li (list of str), table (rows), pagebreak
Spec = list[tuple[str, object]]


def pdf_bytes(spec: Spec, footer: str | None = None, encrypt: str | None = None) -> bytes:
    styles = getSampleStyleSheet()
    flowables = []
    for kind, content in spec:
        if kind == "title":
            flowables.append(Paragraph(str(content), styles["Title"]))
        elif kind == "h1":
            flowables.append(Paragraph(str(content), styles["Heading1"]))
        elif kind == "h2":
            flowables.append(Paragraph(str(content), styles["Heading2"]))
        elif kind == "p":
            flowables.append(Paragraph(str(content), styles["BodyText"]))
        elif kind == "li":
            items = [Paragraph(i, styles["BodyText"]) for i in content]  # type: ignore[union-attr]
            flowables.append(ListFlowable(items, bulletType="bullet"))
        elif kind == "table":
            table = Table(content, hAlign="LEFT")
            table.setStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey)])
            flowables.append(table)
        elif kind == "pagebreak":
            flowables.append(PageBreak())

    def draw_footer(pdf_canvas, doc) -> None:  # noqa: ANN001
        if footer:
            pdf_canvas.setFont("Helvetica", 8)
            pdf_canvas.drawString(50, 25, f"{footer} | Page {doc.page}")

    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, encrypt=encrypt).build(
        flowables, onFirstPage=draw_footer, onLaterPages=draw_footer
    )
    return buffer.getvalue()


def image_only_pdf_bytes(pages: int = 1) -> bytes:
    """Imitates a scanned document: every page is a picture, with no text layer."""
    picture = io.BytesIO()
    Image.new("RGB", (200, 300), "white").save(picture, "PNG")
    buffer = io.BytesIO()
    pdf_canvas = canvas.Canvas(buffer)
    for _ in range(pages):
        picture.seek(0)
        pdf_canvas.drawImage(ImageReader(picture), 50, 50, 400, 600)
        pdf_canvas.showPage()
    pdf_canvas.save()
    return buffer.getvalue()


def blank_pdf_bytes() -> bytes:
    buffer = io.BytesIO()
    pdf_canvas = canvas.Canvas(buffer)
    pdf_canvas.showPage()
    pdf_canvas.save()
    return buffer.getvalue()


def docx_bytes(spec: Spec, bold_paragraphs: list[str] | None = None) -> bytes:
    document = Document()
    for kind, content in spec:
        if kind == "title":
            document.add_heading(str(content), level=0)
        elif kind == "h1":
            document.add_heading(str(content), level=1)
        elif kind == "h2":
            document.add_heading(str(content), level=2)
        elif kind == "p":
            document.add_paragraph(str(content))
        elif kind == "li":
            for item in content:  # type: ignore[union-attr]
                document.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            rows = content  # type: ignore[assignment]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
    for text in bold_paragraphs or []:
        document.add_paragraph().add_run(text).bold = True
    document.sections[0].footer.paragraphs[0].text = "Running footer that must be ignored"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
