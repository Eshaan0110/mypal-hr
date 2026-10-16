"""Generate the FICTIONAL sample policies in sample_data/policies/ (PDF and DOCX).

These exist only so ingestion, retrieval and the admin workflow can be built and tested before
HR provides real, approved policies. Every document says so on its first line. "Demo Health
Services" is an invented organisation; no content here is a policy of The Manipal Group.

Run from backend/:  uv run python scripts/build_sample_policies.py
"""

from pathlib import Path

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import ListFlowable, Paragraph, SimpleDocTemplate, Spacer, Table

OUT_DIR = Path(__file__).resolve().parents[2] / "sample_data" / "policies"
BANNER = (
    "FICTIONAL SAMPLE FOR SOFTWARE DEVELOPMENT ONLY. This document is not a policy of "
    "The Manipal Group or of any real organisation."
)
ORG = "Demo Health Services (fictional)"

# A document is a list of (kind, content): h1/h2 headings, p paragraphs, li bullet lists,
# table rows (first row is the header).
Block = tuple[str, object]


def leave_policy(version: str) -> list[Block]:
    v2 = version == "v2"
    earned_days = 20 if v2 else 18
    notice_days = 10 if v2 else 7
    effective = "1 April 2026" if v2 else "1 January 2026"
    return [
        ("p", BANNER),
        ("title", f"Leave Policy ({version})"),
        ("p", f"Organisation: {ORG}. Policy owner: HR Operations. Effective from {effective}."),
        ("h1", "1. Purpose"),
        (
            "p",
            "This policy explains the types of leave available to employees, how leave is "
            "accrued, and how to apply for it.",
        ),
        ("h1", "2. Scope"),
        (
            "p",
            "This policy applies to permanent employees at all campuses of the organisation. "
            "Trainees are covered only by section 6.",
        ),
        ("h1", "3. Types of leave"),
        ("h2", "3.1 Casual leave"),
        (
            "p",
            "Employees are entitled to 8 days of casual leave per calendar year. Casual leave "
            "may not exceed 3 consecutive working days and cannot be carried forward to the "
            "next year.",
        ),
        ("h2", "3.2 Sick leave"),
        (
            "p",
            "Employees are entitled to 10 days of sick leave per calendar year. A medical "
            "certificate is required for sick leave of more than 2 consecutive days.",
        ),
        ("h2", "3.3 Earned leave"),
        (
            "p",
            f"Employees earn {earned_days} days of earned leave per year, credited monthly. "
            "Unused earned leave may be carried forward up to a maximum of 30 days. Earned "
            "leave is available only after successful completion of probation.",
        ),
        ("h1", "4. How to apply for leave"),
        (
            "li",
            [
                f"Apply in the HR portal at least {notice_days} calendar days in advance for "
                "earned leave.",
                "Casual and sick leave may be applied for on the day of absence or earlier.",
                "The reporting manager approves or rejects the application.",
                "Approved leave appears in the employee's leave record within 2 working days. "
                "This is an estimated processing time, not a policy deadline.",
            ],
        ),
        ("h1", "5. Leave during notice period"),
        (
            "p",
            "Earned leave may not be taken during the notice period without written approval "
            "from the HR Business Partner.",
        ),
        ("h1", "6. Probation and trainees"),
        (
            "p",
            "Employees on probation may take casual and sick leave but not earned leave. "
            "Trainees are entitled to 6 days of casual leave per year and no other leave types.",
        ),
        ("h1", "7. Questions"),
        (
            "p",
            "Questions about this policy, or cases it does not cover, must be referred to the "
            "designated HR contact.",
        ),
    ]


def travel_policy() -> list[Block]:
    return [
        ("p", BANNER),
        ("title", "Travel and Reimbursement Policy (v1)"),
        (
            "p",
            f"Organisation: {ORG}. Policy owner: Finance and HR. Effective from 1 February 2026.",
        ),
        ("h1", "1. Purpose"),
        (
            "p",
            "This policy sets out allowances for official travel and the process for claiming "
            "reimbursement of travel expenses.",
        ),
        ("h1", "2. Daily allowance"),
        ("p", "The daily allowance depends on the employee's grade and the destination city."),
        (
            "table",
            [
                ["Grade", "Metro city (per day)", "Other city (per day)"],
                ["Grade A", "INR 4,000", "INR 3,000"],
                ["Grade B", "INR 3,000", "INR 2,200"],
                ["Grade C", "INR 2,000", "INR 1,500"],
            ],
        ),
        ("h1", "3. Approval before travel"),
        (
            "p",
            "All official travel requires a travel request approved by the reporting manager "
            "before tickets are booked. International travel also requires approval from the "
            "department head.",
        ),
        ("h1", "4. Submitting a claim"),
        ("h2", "4.1 Deadline"),
        (
            "p",
            "Claims must be submitted within 15 calendar days of returning from travel. Claims "
            "submitted later are not reimbursed unless the department head approves an exception.",
        ),
        ("h2", "4.2 Documents required"),
        (
            "li",
            [
                "The approved travel request.",
                "Boarding passes or tickets for all journeys.",
                "Original hotel invoices showing the employee's name.",
                "Receipts for local conveyance above INR 500.",
            ],
        ),
        ("h2", "4.3 Approval and payment"),
        (
            "p",
            "Claims are approved by the reporting manager and then verified by Finance. Payment "
            "is usually made within 10 working days of verification; this is an estimate, not "
            "a guarantee.",
        ),
    ]


def salary_advance_policy() -> list[Block]:
    return [
        ("p", BANNER),
        ("title", "Salary Advance Policy (v1)"),
        ("p", f"Organisation: {ORG}. Policy owner: HR Operations. Effective from 1 March 2026."),
        ("h1", "1. Eligibility"),
        (
            "p",
            "Permanent employees who have completed probation may request a salary advance. "
            "Employees on probation, trainees and employees serving notice are not eligible.",
        ),
        ("h1", "2. Amount"),
        (
            "p",
            "The maximum advance is two months of basic salary. Only one advance may be "
            "outstanding at any time.",
        ),
        ("h1", "3. Repayment"),
        (
            "p",
            "The advance is repaid through equal monthly deductions from salary over a maximum "
            "of 6 instalments, starting from the month after the advance is paid. No interest "
            "is charged.",
        ),
        ("h1", "4. How to apply"),
        (
            "li",
            [
                "Complete form SA-1 (fictional) and state the reason for the request.",
                "The reporting manager recommends the request.",
                "The HR Head approves or rejects it.",
            ],
        ),
        ("h1", "5. Exceptions"),
        (
            "p",
            "Any exception to this policy, including a larger amount or longer repayment "
            "period, requires written approval from the HR Head and is decided case by case.",
        ),
    ]


def write_pdf(blocks: list[Block], path: Path, footer: str) -> None:
    styles = getSampleStyleSheet()
    story = []
    for kind, content in blocks:
        if kind == "title":
            story.append(Paragraph(str(content), styles["Title"]))
        elif kind == "h1":
            story.append(Paragraph(str(content), styles["Heading1"]))
        elif kind == "h2":
            story.append(Paragraph(str(content), styles["Heading2"]))
        elif kind == "p":
            story.append(Paragraph(str(content), styles["BodyText"]))
        elif kind == "li":
            items = [Paragraph(item, styles["BodyText"]) for item in content]  # type: ignore[union-attr]
            story.append(ListFlowable(items, bulletType="bullet"))
        elif kind == "table":
            table = Table(content, hAlign="LEFT")
            table.setStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey)])
            story.append(table)
        story.append(Spacer(1, 3 * mm))

    def draw_footer(canvas, doc) -> None:  # noqa: ANN001 - reportlab callback signature
        canvas.setFont("Helvetica", 8)
        canvas.drawString(20 * mm, 10 * mm, f"{footer} | Page {doc.page}")

    SimpleDocTemplate(str(path), pagesize=A4, title=footer).build(
        story, onFirstPage=draw_footer, onLaterPages=draw_footer
    )


def write_docx(blocks: list[Block], path: Path) -> None:
    document = Document()
    for kind, content in blocks:
        if kind == "title":
            document.add_heading(str(content), level=0)
        elif kind == "h1":
            document.add_heading(str(content), level=1)
        elif kind == "h2":
            document.add_heading(str(content), level=2)
        elif kind == "p":
            document.add_paragraph(str(content))
        elif kind == "li":
            for item in content:  # type: ignore[union-attr]
                document.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            rows = content  # type: ignore[assignment]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
    document.core_properties.title = path.stem
    document.save(str(path))


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    footer = "FICTIONAL SAMPLE - Demo Health Services"
    write_pdf(leave_policy("v1"), OUT_DIR / "leave-policy-v1.pdf", footer)
    write_pdf(leave_policy("v2"), OUT_DIR / "leave-policy-v2.pdf", footer)
    write_docx(travel_policy(), OUT_DIR / "travel-reimbursement-policy-v1.docx")
    write_pdf(salary_advance_policy(), OUT_DIR / "salary-advance-policy-v1.pdf", footer)
    write_docx(salary_advance_policy(), OUT_DIR / "salary-advance-policy-v1.docx")
    for path in sorted(OUT_DIR.iterdir()):
        print(f"wrote {path.relative_to(OUT_DIR.parents[1])}")


if __name__ == "__main__":
    main()
